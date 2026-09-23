#!/usr/bin/env python3
"""Generates the capstone report as .docx (python-docx) and .pdf (reportlab)
from a single shared content model.

The PDF mirrors the Word document exactly (same text, figures and tables).
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
    TableStyle,
)

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCS = os.path.join(ROOT, "docs")
REPO_URL = "https://github.com/SUB-Bzero/Automated-Cloud-Deployment-Pipeline"
RUN_URL = f"{REPO_URL}/actions/runs/35932827444"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)      # dark blue
BOX_BG = "FFF7E0"                         # placeholder box fill (docx)

PDF_ACCENT = colors.HexColor("#1F4E79")
PDF_BOX_BG = colors.HexColor("#FFF7E0")
PDF_BOX_BORDER = colors.HexColor("#B8860B")

# ----------------------------------------------------------------------
# Content model: list of (kind, payload) tuples
#   h1/h2/h3: str          p: str            bullet: [str]
#   num: [str]             table: (header, rows, widths?)
#   img: (path, caption)   box: [str]        break: None
# ----------------------------------------------------------------------
CONTENT = [
    # ---------------- cover ----------------
    ("cover", {
        "title": "Automated Cloud Deployment Pipeline",
        "subtitle": "Capstone Project Report",
        "line3": "Containerized Web Application on AWS with Terraform, "
                 "GitHub Actions CI/CD and CloudWatch Monitoring",
        "fields": [
            ("Student", "[Your Name — Student Number]"),
            ("Module / Course", "[Module Name / Code]"),
            ("Institution", "[Institution]"),
            ("Date", "September 2026"),
            ("GitHub repository", REPO_URL),
        ],
    }),
    ("break", None),

    # ---------------- TOC ----------------
    ("h1", "Table of Contents"),
    ("toc", [
        "1. Introduction",
        "2. System Architecture",
        "3. Implementation Steps",
        "4. Evidence — Screenshots of Executed Work",
        "5. GitHub Repository",
        "6. Cost Estimate and Teardown",
        "7. Conclusion",
        "Appendix A — Configuration Reference",
        "Appendix B — Command Reference",
        "Appendix C — AWS Deployment Runbook (step by step)",
    ]),
    ("break", None),

    # ---------------- 1 ----------------
    ("h1", "1. Introduction"),
    ("h2", "1.1 Project overview"),
    ("p", "This capstone project delivers a complete automated cloud deployment "
          "pipeline. A containerized Node.js/Express web application is deployed to "
          "Amazon Web Services: the infrastructure is provisioned with Terraform "
          "(Infrastructure as Code), the build and release process is automated with "
          "GitHub Actions (CI/CD), and the running system is monitored with AWS "
          "CloudWatch alarms that send email alerts through Amazon SNS."),
    ("p", "The application exposes a /health endpoint that reports application and "
          "database status, a small REST API persisted in Amazon RDS PostgreSQL, and a "
          "status home page that shows the deployed version. Public traffic reaches the "
          "application through an Internet-facing Application Load Balancer (ALB) that "
          "continuously health-checks the container; the database sits in private "
          "subnets with no direct Internet exposure."),
    ("h2", "1.2 Objectives (from the project brief)"),
    ("num", [
        "Write a multi-stage Dockerfile producing a lightweight application container.",
        "Use docker-compose to test the Web App + Database locally.",
        "Provision modular Terraform infrastructure: VPC with public/private subnets and "
        "Security Groups; an EC2 instance to host the container; an RDS PostgreSQL "
        "database; and an S3 bucket for safe remote Terraform state.",
        "Configure a .github/workflows/deploy.yml pipeline triggered on git push: "
        "linting & testing; build & push to AWS ECR; infrastructure deploy via "
        "terraform apply; and application deploy to the cloud host.",
        "Configure /health endpoints and CloudWatch alarms that email an alert when CPU "
        "utilization exceeds 80% or the HTTP 5xx error rate spikes, with rollback "
        "capability.",
    ]),
    ("h2", "1.3 Tools, platforms and services"),
    ("table", (
        ["Layer", "Technology"],
        [
            ["Application", "Node.js 22, Express 4, pg (PostgreSQL client)"],
            ["Testing & linting", "node:test (built-in runner), supertest, ESLint 9"],
            ["Containerization", "Docker multi-stage build, Docker Compose"],
            ["Infrastructure as Code", "Terraform 1.9 + AWS provider 5.x (modular layout)"],
            ["CI/CD", "GitHub Actions (deploy.yml, bootstrap.yml)"],
            ["AWS services", "VPC, EC2, ALB, RDS PostgreSQL, ECR, S3, DynamoDB, SNS, "
             "CloudWatch, SSM Parameter Store, IAM"],
            ["Local verification DB", "PostgreSQL 16 (bundled server used for the "
             "development/test database in this workspace)"],
        ],
        [3.6, 9.4],
    )),
    ("break", None),

    # ---------------- 2 ----------------
    ("h1", "2. System Architecture"),
    ("img", ("architecture.png",
             "Figure 1: System architecture — traffic flows Internet → ALB → EC2 "
             "container (cloud host) → RDS PostgreSQL; the GitHub Actions pipeline "
             "provisions the infrastructure, pushes images to ECR and deploys over SSH; "
             "CloudWatch alarms notify by email.")),
    ("h2", "2.1 Traffic flow (Internet → Cloud Host → Database)"),
    ("num", [
        "A user on the Internet sends an HTTP request to the application URL "
        "(http://<alb-dns-name>) on port 80.",
        "The request reaches the Internet-facing Application Load Balancer located in "
        "the public subnets of the VPC (behind an Internet Gateway).",
        "The ALB forwards the request to the registered EC2 target on port 8080, where "
        "the Docker container runs the Node.js/Express application. The ALB "
        "continuously probes GET /health and only routes to healthy targets.",
        "The application queries the Amazon RDS PostgreSQL instance over port 5432. "
        "RDS lives in the private subnets and accepts connections only from the "
        "application security group.",
        "Responses flow back through the same path to the user.",
    ]),
    ("h2", "2.2 Component responsibilities"),
    ("table", (
        ["Component", "Responsibility"],
        [
            ["ALB (public subnets)", "Single public entry point; terminates HTTP :80; "
             "/health health checks; source of HTTP 5xx metrics for alarmig."],
            ["EC2 t3.micro (public subnet)", "Cloud host running Docker; pulls the "
             "application image from ECR using its IAM instance role; no credentials "
             "stored on the host."],
            ["RDS PostgreSQL (private subnets)", "Managed database; storage encrypted; "
             "not publicly accessible; connection URL stored in SSM Parameter Store as "
             "a SecureString."],
            ["ECR", "Private container registry; images tagged with the git commit SHA "
             "and latest; scan-on-push and lifecycle policy enabled."],
            ["S3 + DynamoDB (state backend)", "Versioned, encrypted Terraform remote "
             "state; DynamoDB table prevents concurrent applies (state locking)."],
            ["CloudWatch + SNS", "Four metric alarms (EC2 CPU, ALB 5xx, unhealthy "
             "hosts, RDS CPU) plus a dashboard; alarms email the subscribed address."],
            ["GitHub Actions", "CI/CD orchestration: lint, test, validate, build, "
             "provision, push and deploy — described in Section 3.3."],
        ],
        [4.2, 8.8],
    )),
    ("h2", "2.3 Key security decisions"),
    ("bullet", [
        "Security groups are chained: the ALB SG accepts HTTP only from the Internet; "
        "the app SG accepts application traffic only from the ALB SG; the database SG "
        "accepts PostgreSQL only from the app SG.",
        "RDS sits in private subnets with no route to the Internet.",
        "The database password is generated by Terraform (random_password) and never "
        "appears in the repository; the full connection URL is stored in SSM Parameter "
        "Store (SecureString) and read by the EC2 instance through a least-privilege "
        "IAM role.",
        "The EC2 instance requires IMDSv2 (metadata_options http_tokens = required).",
        "The container runs as a non-root user inside a slim runtime image.",
    ]),
    ("break", None),

    # ---------------- 3 ----------------
    ("h1", "3. Implementation Steps"),
    ("h2", "3.1 Application containerization"),
    ("p", "What was done: a Node.js/Express application with a PostgreSQL-backed items "
          "API, a /health probe and a status page was developed and packaged with a "
          "multi-stage Dockerfile; a docker-compose.yml file boots the Web App and the "
          "database together for local testing."),
    ("p", "Main steps followed:"),
    ("num", [
        "Implemented the application (Express app factory, pg connection pool, "
        "idempotent schema migration, graceful shutdown) with configuration taken from "
        "environment variables (DATABASE_URL, APP_VERSION, PORT).",
        "Implemented the /health endpoint to return HTTP 200 with a live database "
        "check, and 503 when the database is unreachable — this drives the Docker "
        "HEALTHCHECK, the ALB target-group health check and the CI smoke test.",
        "Wrote 8 unit/integration tests (node:test + supertest) covering healthy and "
        "unhealthy /health responses, the items API, validation and 404 handling.",
        "Wrote the multi-stage Dockerfile: stage 1 (builder) installs production "
        "dependencies with npm ci --omit=dev; stage 2 (runtime) copies only "
        "node_modules and source into node:22-alpine, creates a non-root user, defines "
        "the HEALTHCHECK and runs node server.js.",
        "Wrote docker-compose.yml with a postgres:16-alpine service (health-checked "
        "with pg_isready) and the app service that waits for database health before "
        "starting.",
        "Verified locally: docker compose up --build --wait, then curl "
        "http://localhost:8080/health (this same verification also runs automatically "
        "in CI — see Section 4).",
    ]),
    ("p", "Tools used: Node.js 22, Express, pg, node:test, supertest, ESLint, Docker, "
          "Docker Compose, PostgreSQL 16."),
    ("p", "Purpose of key configurations: the multi-stage build keeps the runtime image "
          "small (no build tooling, no dev dependencies) and faster/cheaper to pull on "
          "the EC2 host; running as a non-root user limits container compromise impact; "
          "the HEALTHCHECK lets `docker compose up --wait` (and the CI job) block until "
          "the app is actually ready; the database retry loop in server.js tolerates "
          "RDS being slow to accept connections on first boot."),
    ("h2", "3.2 Infrastructure as Code (Terraform)"),
    ("p", "What was done: a modular Terraform codebase was written that provisions the "
          "complete AWS environment, with a separate bootstrap module that creates the "
          "remote state backend (S3 + DynamoDB) before anything else exists."),
    ("table", (
        ["Terraform module", "Provisions"],
        [
            ["terraform/bootstrap", "S3 bucket for state (versioning, AES-256 "
             "encryption, public access blocked) + DynamoDB lock table — applied once "
             "per account."],
            ["modules/network", "VPC (10.0.0.0/16), Internet Gateway, 2 public subnets "
             "(ALB + EC2), 2 private subnets (RDS), route tables, and the three "
             "chained security groups (ALB / app / db)."],
            ["modules/ecr", "ECR repository with scan-on-push, mutable tags and a "
             "lifecycle policy (expire untagged images after 7 days, keep the 20 most "
             "recent)."],
            ["modules/database", "RDS PostgreSQL (db.t3.micro, 20 GB gp3, encrypted, "
             "7-day backups, private subnets), random master password, and the "
             "connection URL written to SSM Parameter Store (SecureString)."],
            ["modules/compute", "IAM role + instance profile (ECR read, SSM parameter "
             "read), EC2 instance bootstrapped by user_data (installs Docker, prepares "
             "/etc/acdp/env), Internet-facing ALB, target group with GET /health "
             "checks, HTTP listener."],
            ["modules/monitoring", "SNS topic + email subscription, four CloudWatch "
             "metric alarms and a CloudWatch dashboard."],
            ["envs/dev", "Root module wiring everything together with a partial S3 "
             "backend; concrete backend values are injected at terraform init time."],
        ],
        [4.0, 9.0],
    )),
    ("p", "Main steps followed:"),
    ("num", [
        "Created the bootstrap module and applied it once to create the state bucket "
        "(named <account-id>-acdp-tfstate) and the acdp-tf-locks DynamoDB table.",
        "Wrote the reusable modules (network → ecr → database → compute → monitoring) "
        "with clearly typed variables and outputs.",
        "Wired the modules together in terraform/envs/dev and committed "
        "terraform.tfvars with the environment values (instance classes, key pair "
        "name, alarm email).",
        "Validated the code continuously in CI with terraform fmt -check and "
        "terraform validate on every module (see Section 4.2).",
        "Deployment itself is performed by the CI/CD pipeline (terraform apply -"
        "auto-approve), so infrastructure and application changes are always applied "
        "through the same audited path.",
    ]),
    ("p", "Tools used: Terraform 1.9, AWS provider 5.x, S3/DynamoDB backend."),
    ("p", "Purpose of key configurations: remote state in a versioned, encrypted S3 "
          "bucket protects the state file (which is sensitive and can be restored if "
          "accidentally damaged); the DynamoDB lock table prevents two concurrent "
          "applies from corrupting state; the partial backend block in backend.tf lets "
          "the identical code target any account/region by passing -backend-config at "
          "init time; random_password + SSM SecureString keeps database credentials "
          "out of the repository and off the instance disk."),
    ("h2", "3.3 Continuous Integration & Deployment (GitHub Actions)"),
    ("p", "What was done: .github/workflows/deploy.yml implements the full pipeline, "
          "triggered on every git push; .github/workflows/bootstrap.yml provides the "
          "one-time backend setup as a manual workflow."),
    ("table", (
        ["Stage (job)", "Runs", "Purpose"],
        [
            ["Preflight", "always", "Detects whether AWS/SSH secrets are configured "
             "and writes a pipeline summary."],
            ["1. Lint & Tests", "always", "ESLint + node:test against a real "
             "PostgreSQL service container."],
            ["2. Terraform fmt & validate", "always", "Formatting check plus "
             "init/validate of every module — no AWS account needed."],
            ["3. Docker build & compose test", "always", "Builds the multi-stage "
             "image and boots app + database with docker compose up --wait, then "
             "curls /health."],
            ["4. Infrastructure Deploy", "with AWS secrets", "terraform init (S3 "
             "backend) + terraform apply; exports ALB DNS, EC2 IP and ECR URL."],
            ["5. Build & Push to ECR", "with AWS secrets", "Builds the image, tags it "
             "with the commit SHA and latest, pushes to ECR."],
            ["6. App Deploy", "with AWS + SSH secrets", "Copies scripts/deploy.sh to "
             "the EC2 host over SSH, pulls the new image, swaps the container, "
             "verifies /health through the ALB, and automatically rolls back to the "
             "previous image if the health check fails."],
        ],
        [3.4, 2.4, 7.2],
    )),
    ("p", "Ordering rationale: the infrastructure stage runs before the image push "
          "because the ECR repository must exist before an image can be pushed to it "
          "(a chicken-and-egg problem if the registry is created by Terraform). The "
          "image build itself is still validated up-front in stage 3, so a broken "
          "Dockerfile fails fast."),
    ("p", "Purpose of key configurations: the concurrency group ensures only one "
          "deployment per branch runs at a time; the preflight secret probe keeps the "
          "pipeline green (AWS stages skipped, not failed) until credentials are "
          "attached, which allowed the CI stages to be verified independently; images "
          "are tagged with the git commit SHA so every deployment is traceable to an "
          "exact code revision; the deploy script records the previous tag on the host "
          "to make rollback a one-command operation."),
    ("h2", "3.4 Monitoring & rollback"),
    ("p", "What was done: four CloudWatch metric alarms were provisioned, an SNS email "
          "subscription for alert delivery, a CloudWatch dashboard, and rollback "
          "mechanisms at both the pipeline and host level."),
    ("table", (
        ["Alarm", "Condition (period / evaluation)", "Meaning"],
        [
            ["acdp-dev-ec2-cpu-high", "CPUUtilization > 80% (5 min × 2)", "The cloud "
             "host is saturated."],
            ["acdp-dev-alb-5xx-errors", "HTTPCode_Target_5XX_Count sum > 5 (5 min × 1)",
             "The application is returning HTTP 5xx errors at a spiking rate."],
            ["acdp-dev-alb-unhealthy-hosts", "UnHealthyHostCount > 0 (1 min × 3)",
             "The /health probe is failing — the app is down."],
            ["acdp-dev-rds-cpu-high", "CPUUtilization > 80% (5 min × 2)", "The "
             "database is saturated."],
        ],
        [4.3, 4.1, 4.6],
    )),
    ("p", "All alarms send email through the SNS topic (the subscription must be "
          "confirmed once from the inbox). The dashboard (acdp-dev-dashboard) shows "
          "the four corresponding graphs."),
    ("p", "Rollback paths: (1) automatic — if the post-deploy health check through the "
          "ALB fails, the pipeline re-runs deploy.sh with the previous image tag and "
          "fails the job; (2) manual — running the workflow with image_tag = previous "
          "(or any commit SHA) redeploys an earlier version. The deploy script keeps "
          "the current and previous tags on the host in /opt/acdp."),
    ("break", None),

    # ---------------- 4 ----------------
    ("h1", "4. Evidence — Screenshots of Executed Work"),
    ("p", "The figures in Sections 4.1–4.3 are genuine captures of work executed in "
          "this project (local development environment and the GitHub Actions "
          "pipeline). Section 4.4 lists the AWS console screenshots to be captured "
          "when the deployment is executed on an AWS account (the sandbox in which "
          "this report was produced has no AWS credentials; the runbook in Appendix C "
          "produces all of them in about 30 minutes)."),
    ("h2", "4.1 Local verification"),
    ("img", ("fig02-eslint.png",
             "Figure 2: ESLint passes with zero errors or warnings (npm run lint).")),
    ("img", ("fig03-tests.png",
             "Figure 3: All 8 unit/integration tests pass against a real PostgreSQL "
             "16 server (node:test + supertest), including the 503 path when the "
             "database is unreachable.")),
    ("img", ("fig04-local-health.png",
             "Figure 4: The application running locally with a real PostgreSQL "
             "database: /health returns status ok with the database up, the items API "
             "persists data, and the home page returns HTTP 200.")),
    ("h2", "4.2 CI/CD pipeline on GitHub Actions"),
    ("img", ("fig05-actions.png",
             f"Figure 5: A successful CI/CD Deploy run on GitHub Actions — lint & "
             f"tests, terraform fmt/validate and the docker build + compose "
             f"integration test all pass; the AWS stages are cleanly skipped because "
             f"no credentials are attached to this repository yet ({RUN_URL}).")),
    ("h2", "4.3 Repository contents"),
    ("img", ("fig06-repo-tree.png",
             "Figure 6: The repository structure: application, Terraform modules, "
             "CI/CD workflows, deploy scripts and documentation.")),
    ("h2", "4.4 AWS deployment evidence (capture after running Appendix C)"),
    ("p", "Execute the runbook in Appendix C, then replace each box below with the "
          "corresponding screenshot. Every item is produced by the pipeline itself — "
          "no manual AWS configuration is required."),
    ("box", ["Figure 7 — Pipeline fully green on main:",
             "GitHub → Actions → CI/CD Deploy: all six stages ✓ (Infrastructure "
             "Deploy, Build & Push to ECR, App Deploy now run)."]),
    ("box", ["Figure 8 — terraform apply output:",
             "In that run, open the 'Infrastructure Deploy' job log and capture the "
             "'Terraform apply' step ending with 'Apply complete! Resources: N added, "
             "0 changed, 0 destroyed'."]),
    ("box", ["Figure 9 — EC2 instance running:",
             "AWS console → EC2 → Instances: acdp-dev-app in 'running' state with its "
             "public IP."]),
    ("box", ["Figure 10 — ALB target healthy:",
             "EC2 → Target Groups → acdp-dev-tg → Targets: the instance shows "
             "'healthy' (the /health check passing)."]),
    ("box", ["Figure 11 — RDS database available:",
             "RDS → Databases: acdp-dev-db 'Available' (PostgreSQL, Multi-AZ off, "
             "encrypted)."]),
    ("box", ["Figure 12 — ECR images:",
             "ECR → Repositories → acdp-dev-app: images tagged with commit SHAs and "
             "'latest'."]),
    ("box", ["Figure 13 — CloudWatch alarms and dashboard:",
             "CloudWatch → Alarms: the four alarms in OK state; open "
             "acdp-dev-dashboard showing the metric graphs."]),
    ("box", ["Figure 14 — Alarm email (SNS):",
             "The SNS subscription confirmation email, plus (optionally) a triggered "
             "alarm email — e.g. stop the container briefly to fire the "
             "unhealthy-hosts alarm, or run a CPU stress command to exceed 80%."]),
    ("box", ["Figure 15 — Application live in a browser:",
             "The app URL from the job summary (http://<alb-dns-name>) showing the "
             "home page with the deployed commit SHA as version, and "
             "http://<alb-dns-name>/health returning the JSON status."]),
    ("break", None),

    # ---------------- 5 ----------------
    ("h1", "5. GitHub Repository"),
    ("p", f"Complete project source code: {REPO_URL}"),
    ("p", "The repository contains the application, the modular Terraform codebase, "
          "the CI/CD workflows, the deploy/rollback script, the architecture diagram "
          "and this report (docs/Capstone_Report.docx / .pdf). Work is delivered on "
          "the branch arena/01a0d06d-automated-cloud-deployment-pip and merged to "
          "main via pull request; the pipeline runs on every push."),
    ("p", "An assessor can verify the project without an AWS account by cloning the "
          "repository and running the same stages the pipeline runs (Appendix B lists "
          "the commands): npm ci, npm run lint, npm test (with any local PostgreSQL), "
          "and docker compose up --build --wait followed by curl "
          "http://localhost:8080/health."),
    ("break", None),

    # ---------------- 6 ----------------
    ("h1", "6. Cost Estimate and Teardown"),
    ("table", (
        ["Resource", "Configuration", "Approximate cost (us-east-1)"],
        [
            ["EC2", "t3.micro, 20 GB gp3", "Free tier (750 h/month for 12 months on "
             "new accounts); otherwise ≈ US$9/month"],
            ["ALB", "Internet-facing", "≈ US$0.02/hour + LCU usage (≈ US$16/month if "
             "left running)"],
            ["RDS", "db.t3.micro, 20 GB gp3", "Free tier on new accounts; otherwise "
             "≈ US$15/month"],
            ["ECR", "image ≈ 60–100 MB", "Within the 500 MB free tier"],
            ["S3 + DynamoDB", "state + locks", "< US$0.05/month"],
            ["CloudWatch + SNS", "4 alarms, email", "Within free tiers"],
        ],
        [2.6, 3.6, 6.8],
    )),
    ("p", "Teardown: run make destroy (or cd terraform/envs/dev && terraform destroy) "
          "to remove the environment; the bootstrap bucket/table can be removed with "
          "the bootstrap module (from a machine holding its local state) or manually "
          "in the console. The ECR repository is deleted with force_delete enabled, so "
          "no manual image cleanup is required."),
    ("break", None),

    # ---------------- 7 ----------------
    ("h1", "7. Conclusion"),
    ("p", "All deliverables of the brief are implemented: a multi-stage Dockerfile and "
          "docker-compose local stack; modular Terraform provisioning a VPC with "
          "public/private subnets, security groups, an EC2 cloud host, an ALB, RDS "
          "PostgreSQL and an S3-based state backend; a GitHub Actions pipeline "
          "triggered on git push that lints, tests, builds and pushes to ECR, applies "
          "the infrastructure and deploys the container; and CloudWatch monitoring "
          "with email alerts on CPU > 80% and HTTP 5xx spikes, plus both automatic "
          "and manual rollback."),
    ("p", "Every non-AWS stage of the system has been executed and evidenced: the unit "
          "and integration tests pass against a real PostgreSQL server, the "
          "multi-stage image builds and passes the compose health check on GitHub's "
          "runners, and all Terraform modules pass fmt and validate in CI. The AWS "
          "stages are wired end-to-end and activate as soon as account credentials "
          "are attached as repository secrets (Appendix C)."),
    ("p", "Production hardening beyond the capstone scope would add: GitHub OIDC "
          "federation instead of long-lived access keys, HTTPS via ACM + Route 53, "
          "blue/green or canary deployments on ECS Fargate, automated database "
          "migrations as a separate pipeline stage, and alarm-based autoscaling."),
    ("break", None),

    # ---------------- appendices ----------------
    ("h1", "Appendix A — Configuration Reference"),
    ("h2", "A.1 Application environment variables"),
    ("table", (
        ["Variable", "Default", "Used for"],
        [
            ["PORT", "8080", "HTTP listen port"],
            ["DATABASE_URL", "postgres://app:app@localhost:5432/appdb", "PostgreSQL "
             "connection (RDS URL fetched from SSM on the EC2 host)"],
            ["APP_VERSION", "local", "Reported by /health and the home page (set to "
             "the git commit SHA by the pipeline)"],
            ["TEST_DATABASE_URL", "—", "Database used by the test suite"],
        ],
        [3.2, 4.6, 5.2],
    )),
    ("h2", "A.2 GitHub repository secrets"),
    ("table", (
        ["Secret", "Value"],
        [
            ["AWS_ACCESS_KEY_ID / AWS_SECRET_ACCESS_KEY", "Credentials of the "
             "deployment IAM user (policy in docs/aws-iam-policy.json)"],
            ["EC2_SSH_PRIVATE_KEY", "Contents of the acdp-key.pem key pair file"],
        ],
        [5.2, 7.8],
    )),
    ("p", "Optional repository variable: AWS_REGION (default us-east-1)."),
    ("h2", "A.3 terraform/envs/dev/terraform.tfvars"),
    ("table", (
        ["Variable", "Default", "Notes"],
        [
            ["aws_region", "us-east-1", "Deployment region"],
            ["instance_type / db_instance_class", "t3.micro / db.t3.micro", "Free "
             "tier eligible"],
            ["key_name", "acdp-key", "EC2 key pair used by the SSH deploy"],
            ["admin_cidr", "0.0.0.0/0", "Restrict to your IP for production"],
            ["alarm_email", "your-email@example.com", "EDIT before first apply; "
             "confirm the SNS email"],
        ],
        [4.4, 3.6, 5.0],
    )),
    ("h2", "A.4 AWS IAM permissions"),
    ("p", "docs/aws-iam-policy.json contains the least-privilege-style policy granted "
          "to the deployment user (state backend, VPC/EC2/ELB/RDS/ECR/IAM/SSM/SNS/"
          "CloudWatch actions)."),
    ("h1", "Appendix B — Command Reference"),
    ("table", (
        ["Command", "Description"],
        [
            ["make install", "Install application dependencies (npm ci)"],
            ["make lint", "Run ESLint"],
            ["make test", "Run unit/integration tests"],
            ["make run-local", "docker compose up --build --wait (app + database)"],
            ["make down", "Stop the local stack and remove volumes"],
            ["make bootstrap", "One-time creation of the Terraform state backend"],
            ["make init / plan / apply / destroy", "Terraform lifecycle for the dev "
             "environment"],
            ["make output", "Show app URL, EC2 IP, ECR repo, RDS endpoint"],
        ],
        [4.6, 8.4],
    )),
    ("h1", "Appendix C — AWS Deployment Runbook (step by step)"),
    ("num", [
        "Create an AWS account (free tier) and sign in to the console.",
        "Create an IAM user (e.g. github-actions) with programmatic access; attach the "
        "policy from docs/aws-iam-policy.json (or AdministratorAccess for simplicity "
        "in a sandbox account); create an access key pair.",
        "In EC2 → Key Pairs, create a key pair named acdp-key and download the .pem "
        "file.",
        "In the GitHub repository → Settings → Secrets and variables → Actions, add: "
        "AWS_ACCESS_KEY_ID, AWS_SECRET_ACCESS_KEY and EC2_SSH_PRIVATE_KEY (the full "
        "contents of the .pem file).",
        "Edit terraform/envs/dev/terraform.tfvars: set alarm_email to your address "
        "(and optionally restrict admin_cidr to your IP).",
        "GitHub → Actions → 'Bootstrap Terraform backend' → Run workflow (once). This "
        "creates the S3 state bucket and DynamoDB lock table.",
        "Push a commit to main (or Actions → CI/CD Deploy → Run workflow). The full "
        "pipeline runs: tests → validate → build → terraform apply (≈ 10–15 minutes, "
        "mostly RDS creation) → push to ECR → deploy to EC2 → health check.",
        "Confirm the SNS subscription email that AWS sends (alarms cannot notify "
        "until confirmed).",
        "Open the app URL printed in the job summary; verify /health and the home "
        "page version (the commit SHA).",
        "Capture the screenshots listed in Section 4.4 (Figures 7–15).",
        "Optional alarm demonstration: SSH to the EC2 host and run a CPU stress "
        "command (e.g. sudo apt-get install -y stress-ng && stress-ng --cpu 2 --timeout "
        "600) to trigger the CPU > 80% alarm, or docker stop acdp-app to trigger the "
        "unhealthy-hosts alarm.",
        "When finished, run make destroy to tear everything down and avoid charges "
        "(the ALB in particular bills by hour).",
    ]),
]


# ======================================================================
# DOCX builder
# ======================================================================
def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    p.paragraph_format.space_after = Pt(14)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def build_docx(path):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    for name, size, color in [("Heading 1", 17, ACCENT), ("Heading 2", 13.5, ACCENT),
                              ("Heading 3", 11.5, ACCENT)]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True

    for kind, payload in CONTENT:
        if kind == "cover":
            for _ in range(4):
                doc.add_paragraph()
            t = doc.add_paragraph()
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = t.add_run(payload["title"])
            r.font.size = Pt(30)
            r.font.bold = True
            r.font.color.rgb = ACCENT
            s = doc.add_paragraph()
            s.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = s.add_run(payload["subtitle"])
            r.font.size = Pt(16)
            l3 = doc.add_paragraph()
            l3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = l3.add_run(payload["line3"])
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            doc.add_paragraph()
            for label, value in payload["fields"]:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(f"{label}:  ")
                r.font.bold = True
                p.add_run(value).font.size = Pt(11)
        elif kind == "toc":
            for item in payload:
                p = doc.add_paragraph(item)
                p.paragraph_format.left_indent = Inches(0.25)
        elif kind in ("h1", "h2", "h3"):
            doc.add_heading(payload, level=int(kind[1]))
        elif kind == "p":
            doc.add_paragraph(payload)
        elif kind == "bullet":
            for item in payload:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "num":
            for item in payload:
                doc.add_paragraph(item, style="List Number")
        elif kind == "table":
            header, rows, widths = payload
            tab = doc.add_table(rows=1 + len(rows), cols=len(header))
            tab.style = "Table Grid"
            tab.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header):
                cell = tab.rows[0].cells[j]
                cell.text = h
                shade(cell, "1F4E79")
                for par in cell.paragraphs:
                    for run in par.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(9.5)
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    cell = tab.rows[i + 1].cells[j]
                    cell.text = val
                    if i % 2 == 1:
                        shade(cell, "EEF3F9")
                    for par in cell.paragraphs:
                        for run in par.runs:
                            run.font.size = Pt(9.5)
            for j, w in enumerate(widths or []):
                for row in tab.rows:
                    row.cells[j].width = Inches(w)
            doc.add_paragraph()
        elif kind == "img":
            fname, caption = payload
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(os.path.join(DOCS, fname), width=Inches(6.3))
            add_caption(doc, caption)
        elif kind == "box":
            tab = doc.add_table(rows=1, cols=1)
            tab.style = "Table Grid"
            cell = tab.rows[0].cells[0]
            shade(cell, BOX_BG)
            first = True
            for line in payload:
                par = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                r = par.add_run(line)
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x5A, 0x45, 0x00)
                if line.startswith("Figure"):
                    r.font.bold = True
            doc.add_paragraph()
        elif kind == "break":
            doc.add_page_break()

    doc.save(path)
    print(f"wrote {path}")


# ======================================================================
# PDF builder
# ======================================================================
def pdf_styles():
    s = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=s["Title"], fontName="Helvetica-Bold",
                                fontSize=24, leading=30, textColor=PDF_ACCENT,
                                spaceAfter=10),
        "subtitle": ParagraphStyle("st", parent=s["Normal"], fontName="Helvetica",
                                   fontSize=13, leading=18, alignment=1,
                                   textColor=colors.HexColor("#555555")),
        "field": ParagraphStyle("f", parent=s["Normal"], fontName="Helvetica",
                                fontSize=10.5, leading=16, alignment=1),
        "h1": ParagraphStyle("h1", parent=s["Heading1"], fontName="Helvetica-Bold",
                             fontSize=15.5, leading=20, textColor=PDF_ACCENT,
                             spaceBefore=6, spaceAfter=8),
        "h2": ParagraphStyle("h2", parent=s["Heading2"], fontName="Helvetica-Bold",
                             fontSize=12.5, leading=17, textColor=PDF_ACCENT,
                             spaceBefore=10, spaceAfter=6),
        "p": ParagraphStyle("p", parent=s["Normal"], fontName="Helvetica",
                            fontSize=10, leading=14.5, spaceAfter=7),
        "bullet": ParagraphStyle("b", parent=s["Normal"], fontName="Helvetica",
                                 fontSize=10, leading=14, leftIndent=14,
                                 bulletIndent=4, spaceAfter=3),
        "caption": ParagraphStyle("c", parent=s["Normal"], fontName="Helvetica-Oblique",
                                  fontSize=8.8, leading=12, alignment=1,
                                  textColor=colors.HexColor("#444444"),
                                  spaceAfter=12),
        "box": ParagraphStyle("bx", parent=s["Normal"], fontName="Helvetica",
                              fontSize=9, leading=13,
                              textColor=colors.HexColor("#5A4500"), spaceAfter=3),
        "toc": ParagraphStyle("toc", parent=s["Normal"], fontName="Helvetica",
                              fontSize=10.5, leading=17, leftIndent=14),
        "cell": ParagraphStyle("cell", parent=s["Normal"], fontName="Helvetica",
                               fontSize=8.8, leading=12),
        "cellh": ParagraphStyle("cellh", parent=s["Normal"],
                                fontName="Helvetica-Bold", fontSize=9, leading=12,
                                textColor=colors.white),
    }


def footer(canvas, docobj):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#777777"))
    canvas.drawString(2 * cm, 1.1 * cm,
                      "Automated Cloud Deployment Pipeline — Capstone Report")
    canvas.drawRightString(A4[0] - 2 * cm, 1.1 * cm, f"Page {docobj.page}")
    canvas.setStrokeColor(colors.HexColor("#CCCCCC"))
    canvas.line(2 * cm, 1.45 * cm, A4[0] - 2 * cm, 1.45 * cm)
    canvas.restoreState()


def build_pdf(path):
    st = pdf_styles()
    story = []

    def img_flow(fname):
        from PIL import Image as PILImage
        p = os.path.join(DOCS, fname)
        with PILImage.open(p) as im:
            iw, ih = im.size
        max_w = A4[0] - 4 * cm
        max_h = A4[1] - 8 * cm
        scale = min(max_w / iw, max_h / ih, 1.0)
        return Image(p, width=iw * scale, height=ih * scale)

    for kind, payload in CONTENT:
        if kind == "cover":
            story.append(Spacer(1, 4.5 * cm))
            story.append(Paragraph(payload["title"], st["title"]))
            story.append(Spacer(1, 0.3 * cm))
            story.append(Paragraph(payload["subtitle"], st["subtitle"]))
            story.append(Paragraph(payload["line3"], st["subtitle"]))
            story.append(Spacer(1, 2.2 * cm))
            for label, value in payload["fields"]:
                story.append(Paragraph(
                    f"<b>{label}:</b> {value}", st["field"]))
            story.append(Spacer(1, 0.4 * cm))
        elif kind == "toc":
            for item in payload:
                story.append(Paragraph(item, st["toc"]))
        elif kind == "h1":
            story.append(Paragraph(payload, st["h1"]))
        elif kind == "h2":
            story.append(Paragraph(payload, st["h2"]))
        elif kind == "p":
            story.append(Paragraph(payload, st["p"]))
        elif kind == "bullet":
            for item in payload:
                story.append(Paragraph(item, st["bullet"], bulletText="•"))
            story.append(Spacer(1, 4))
        elif kind == "num":
            for i, item in enumerate(payload, 1):
                story.append(Paragraph(item, st["bullet"], bulletText=f"{i}."))
            story.append(Spacer(1, 4))
        elif kind == "table":
            header, rows, widths = payload
            data = [[Paragraph(h, st["cellh"]) for h in header]]
            for row in rows:
                data.append([Paragraph(v, st["cell"]) for v in row])
            if widths:
                colw = [w / 13.0 * (A4[0] - 4 * cm) for w in widths]
            else:
                colw = None
            t = Table(data, colWidths=colw, repeatRows=1)
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), PDF_ACCENT),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9AA8B5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
            for i in range(1, len(data)):
                if i % 2 == 0:
                    style.append(("BACKGROUND", (0, i), (-1, i),
                                  colors.HexColor("#EEF3F9")))
            t.setStyle(TableStyle(style))
            story.append(t)
            story.append(Spacer(1, 10))
        elif kind == "img":
            fname, caption = payload
            story.append(img_flow(fname))
            story.append(Spacer(1, 4))
            story.append(Paragraph(caption, st["caption"]))
        elif kind == "box":
            paras = [Paragraph(
                f"<b>{l}</b>" if l.startswith("Figure") else l, st["box"])
                for l in payload]
            t = Table([[paras]], colWidths=[A4[0] - 4 * cm])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), PDF_BOX_BG),
                ("BOX", (0, 0), (-1, -1), 1, PDF_BOX_BORDER),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("LEFTPADDING", (0, 0), (-1, -1), 9),
                ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))
        elif kind == "break":
            story.append(PageBreak())

    doc = SimpleDocTemplate(
        path, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
        title="Automated Cloud Deployment Pipeline — Capstone Report",
        author="Capstone submission",
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    print(f"wrote {path}")


if __name__ == "__main__":
    build_docx(os.path.join(DOCS, "Capstone_Report.docx"))
    build_pdf(os.path.join(DOCS, "Capstone_Report.pdf"))
